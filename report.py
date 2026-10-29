"""
report.py — Attack-surface scan report generation.

build_xlsx(result)        -> bytes   full technical workbook (openpyxl)
build_pdf(result)         -> bytes   executive report (reportlab)
build_research_pdf(item)  -> bytes   Big-4 style threat-intel brief (reportlab)
"""

import io
import re
import html as _html
from datetime import datetime, timezone

# ── openpyxl (XLSX) ──────────────────────────────────────────────────────────
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── reportlab (PDF) ──────────────────────────────────────────────────────────
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak,
)

# ── Shared palette ───────────────────────────────────────────────────────────
SEV_COLORS = {
    "CRITICAL": "C00000", "HIGH": "E8741C", "MEDIUM": "C9A400",
    "LOW": "2E8B57", "INFO": "5673A0",
}
RISK_COLORS = {
    "CRITICAL": "C00000", "HIGH": "E8741C", "MEDIUM": "C9A400",
    "LOW": "2E8B57", "CLEAN": "2E8B57", "UNKNOWN": "5673A0",
}
SEV_ORDER = {"CRITICAL": 5, "HIGH": 4, "MEDIUM": 3, "LOW": 2, "INFO": 1}


def _fmt_ts(ts: str) -> str:
    try:
        return datetime.fromisoformat(ts.replace("Z", "+00:00")).strftime("%Y-%m-%d %H:%M UTC")
    except Exception:
        return ts or datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


def report_filename(result: dict, ext: str) -> str:
    q = re.sub(r"[^\w.-]", "_", (result.get("query", "scan") or "scan"))[:60]
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    return f"attack_surface_{q}_{stamp}.{ext}"


# ════════════════════════════════════════════════════════════════════════════
#  XLSX
# ════════════════════════════════════════════════════════════════════════════

_HDR_FILL = PatternFill("solid", fgColor="0B2545")
_HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
_TITLE_FONT = Font(bold=True, size=14, color="0B2545")
_THIN = Side(style="thin", color="D0D7E2")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = _HDR_FILL
        cell.font = _HDR_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = _BORDER


def _autosize(ws, max_width=70):
    for col in ws.columns:
        length = 0
        letter = None
        for cell in col:
            if letter is None:
                letter = get_column_letter(cell.column)
            try:
                v = str(cell.value) if cell.value is not None else ""
            except Exception:
                v = ""
            length = max(length, len(v))
        if letter:
            ws.column_dimensions[letter].width = min(max_width, max(10, length + 2))


def _sev_fill(sev):
    return PatternFill("solid", fgColor=SEV_COLORS.get(sev, "5673A0"))


def _add_sheet(wb, title, headers, rows, sev_col=None):
    ws = wb.create_sheet(title=title[:31])
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    ws.freeze_panes = "A2"
    for r in rows:
        ws.append(r)
        ridx = ws.max_row
        for c in range(1, len(headers) + 1):
            ws.cell(row=ridx, column=c).border = _BORDER
            ws.cell(row=ridx, column=c).alignment = Alignment(vertical="top", wrap_text=True)
        if sev_col is not None and 0 <= sev_col < len(r):
            sev = str(r[sev_col]).upper()
            cell = ws.cell(row=ridx, column=sev_col + 1)
            cell.fill = _sev_fill(sev)
            cell.font = Font(bold=True, color="FFFFFF")
    _autosize(ws)
    return ws


def build_xlsx(result: dict) -> bytes:
    wb = Workbook()
    rs = result.get("risk_score", {}) or {}
    findings = result.get("findings", []) or []
    sev_counts = {k: 0 for k in SEV_ORDER}
    for f in findings:
        sev_counts[f.get("severity", "INFO")] = sev_counts.get(f.get("severity", "INFO"), 0) + 1

    # ── Summary sheet ────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "JATAYU — External Attack Surface Report"
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:B1")
    meta = [
        ("Target", result.get("query", "")),
        ("Type", (result.get("type", "") or "").upper()),
        ("Scan time", _fmt_ts(result.get("timestamp", ""))),
        ("Risk score", f"{rs.get('score', 0)}/100"),
        ("Risk level", rs.get("level", "UNKNOWN")),
        ("Live hosts", sum(1 for h in result.get("hosts", []) if h.get("alive"))),
        ("Hosts scanned", len(result.get("hosts", []))),
        ("Subdomains found", len(result.get("subdomains", []))),
        ("CVEs found", len(result.get("vulnerabilities", []))),
        ("Total findings", len(findings)),
        ("  Critical", sev_counts["CRITICAL"]),
        ("  High", sev_counts["HIGH"]),
        ("  Medium", sev_counts["MEDIUM"]),
        ("  Low", sev_counts["LOW"]),
        ("  Info", sev_counts["INFO"]),
        ("Email grade", (result.get("email_security", {}) or {}).get("grade", "—")),
        ("Sources used", ", ".join(result.get("sources_used", []))),
    ]
    row = 3
    for k, v in meta:
        ws.cell(row=row, column=1, value=k).font = Font(bold=True)
        ws.cell(row=row, column=2, value=v)
        row += 1
    # Risk factors
    ws.cell(row=row + 1, column=1, value="Risk factors").font = Font(bold=True, size=12)
    row += 2
    for fct in rs.get("factors", []):
        ws.cell(row=row, column=1, value="• " + str(fct))
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 60

    # ── Prioritised remediation ──────────────────────────────────────────────
    _add_sheet(wb, "Prioritized Remediation",
        ["#", "Severity", "Category", "Issue", "Affected", "Recommended Fix"],
        [[i + 1, r.get("severity", ""), r.get("category", ""), r.get("title", ""),
          r.get("affected_count", 0), r.get("remediation", "")]
         for i, r in enumerate(result.get("remediation", []))],
        sev_col=1)

    # ── Hosts ────────────────────────────────────────────────────────────────
    host_rows = []
    for h in result.get("hosts", []):
        fcount = len(h.get("findings", []))
        top = "—"
        if h.get("findings"):
            top = max(h["findings"], key=lambda f: SEV_ORDER.get(f.get("severity", "INFO"), 0)).get("severity", "—")
        host_rows.append([
            h.get("host", ""),
            "LIVE" if h.get("alive") else "resolved",
            h.get("status", "") or "",
            ", ".join(h.get("ips", [])[:3]),
            len(h.get("ports", [])),
            len(h.get("cves", [])),
            fcount,
            top,
            ", ".join(h.get("tech", [])[:5]),
        ])
    _add_sheet(wb, "Hosts",
        ["Host", "State", "HTTP", "IP(s)", "Ports", "CVEs", "Findings", "Top Sev", "Tech"],
        host_rows, sev_col=7)

    # ── Web findings ─────────────────────────────────────────────────────────
    _add_sheet(wb, "Web Findings",
        ["Severity", "Category", "Asset", "Finding", "Evidence", "Remediation"],
        [[f.get("severity", ""), f.get("category", ""), f.get("asset", ""),
          f.get("title", ""), f.get("evidence", ""), f.get("remediation", "")]
         for f in findings],
        sev_col=0)

    # ── Vulnerabilities ──────────────────────────────────────────────────────
    _add_sheet(wb, "Vulnerabilities",
        ["CVE", "NVD Link"],
        [[c, f"https://nvd.nist.gov/vuln/detail/{c}"] for c in result.get("vulnerabilities", [])])

    # ── Open ports ───────────────────────────────────────────────────────────
    port_rows = []
    for ip, pd in (result.get("ports_services", {}) or {}).items():
        port_rows.append([ip, ", ".join(str(p) for p in pd.get("ports", [])),
                          ", ".join(pd.get("hostnames", [])[:3]),
                          ", ".join(pd.get("tags", [])[:6])])
    _add_sheet(wb, "Open Ports", ["IP", "Ports", "Hostnames", "Tags"], port_rows)

    # ── DNS ──────────────────────────────────────────────────────────────────
    dns_rows = []
    for rtype, vals in (result.get("dns", {}) or {}).items():
        for v in vals:
            dns_rows.append([rtype, v])
    _add_sheet(wb, "DNS", ["Type", "Value"], dns_rows)

    # ── Email security ───────────────────────────────────────────────────────
    em = result.get("email_security", {}) or {}
    if em:
        em_rows = [
            ["Grade", em.get("grade", "—")],
            ["Score", f"{em.get('score', 0)}/100"],
            ["SPF", em.get("spf", "—") or "—"],
            ["SPF strict (-all)", "Yes" if em.get("spf_strict") else "No"],
            ["DMARC", em.get("dmarc", "—") or "—"],
            ["DMARC policy", em.get("dmarc_policy", "—") or "—"],
            ["DMARC enforced", "Yes" if em.get("dmarc_enforce") else "No"],
            ["MTA-STS", "Yes" if em.get("mta_sts") else "No"],
        ]
        for iss in em.get("issues", []):
            em_rows.append(["Issue", iss])
        _add_sheet(wb, "Email Security", ["Field", "Value"], em_rows)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ════════════════════════════════════════════════════════════════════════════
#  XLSX — Risk / Findings / Threat-Intel / Remediation / Custom report
# ════════════════════════════════════════════════════════════════════════════

def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _summary_sheet(wb, title: str, rows: list, first: bool = True):
    """A simple key/value summary sheet (reused by all workbook builders)."""
    ws = wb.active if first else wb.create_sheet()
    ws.title = "Summary"
    ws["A1"] = title
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:B1")
    r = 3
    for k, v in rows:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v)
        r += 1
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 60
    return ws


def _stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


def _rem_text(rem: dict) -> str:
    """Flatten a structured remediation dict into a single readable cell."""
    if not rem:
        return ""
    parts = []
    if rem.get("sla"):
        parts.append(f"SLA: {rem['sla']}")
    for i, s in enumerate(rem.get("steps", []), 1):
        parts.append(f"{i}. {s}")
    adv = rem.get("advisories") or []
    if adv:
        parts.append("Advisories: " + " | ".join(a.get("url", "") for a in adv if a.get("url")))
    return "\n".join(parts)


def build_risk_xlsx(risk_result: dict, title: str = "JATAYU — Risk Prioritization") -> bytes:
    """Risk prioritization workbook: Summary + Prioritized Risk + Remediation."""
    items = (risk_result or {}).get("items", []) or []
    summ = (risk_result or {}).get("summary", {}) or {}
    wb = Workbook()
    _summary_sheet(wb, title, [
        ("Generated", _stamp()),
        ("Total findings", summ.get("total_findings", 0)),
        ("Affected assets", summ.get("total_assets", 0)),
        ("Unique CVEs", summ.get("unique_cves", 0)),
        ("Actively exploited (KEV)", summ.get("exploited_cves", 0)),
        ("Ransomware-linked", summ.get("ransomware_cves", 0)),
        ("Critical priorities", summ.get("critical_priorities", 0)),
        ("EPSS ≥ 0.5", summ.get("high_epss", 0)),
        ("In exposure window (AWE > 0)", summ.get("positive_awe_count", 0)),
        ("Avg AWE days", summ.get("avg_awe_days", 0.0)),
        ("Max AWE days", summ.get("max_awe_days", 0.0)),
    ])

    def _pct(v):
        return f"{round(v * 100, 1)}%" if isinstance(v, (int, float)) else ""

    def _awe_fmt(it):
        v = it.get("awe_days")
        if v is None:
            return ""
        return f"+{v}d (exposed)" if v > 0 else f"{v}d (safe)"

    _add_sheet(wb, "Prioritized Risk",
        ["Priority", "Score", "CVE", "Severity", "CVSS", "CVSS Ver", "EPSS",
         "MTTW (days)", "MTTR (days)", "AWE (days)",
         "Exploited (KEV)", "Ransomware", "Attack Vector", "Assets", "Vendor", "Product",
         "Vulnerability", "Recommended Action"],
        [[it.get("priority", ""), it.get("priority_score", ""), it.get("cve_id", ""),
          it.get("severity", ""), it.get("cvss", ""), it.get("cvss_version", ""),
          _pct(it.get("epss")), it.get("mttw_days", ""), it.get("mttr_days", ""),
          _awe_fmt(it), "Yes" if it.get("in_kev") else "No",
          "Yes" if it.get("ransomware") else "No", it.get("attack_vector", ""),
          it.get("affected_assets", 0), it.get("vendor", ""), it.get("product", ""),
          it.get("plugin_name", ""), it.get("action", "")]
         for it in items],
        sev_col=3)

    _add_sheet(wb, "Remediation",
        ["CVE", "Priority", "SLA / Due", "Vendor", "Product", "Detailed Remediation", "Advisories"],
        [[it.get("cve_id", ""), it.get("priority", ""),
          (it.get("remediation") or {}).get("sla", ""), it.get("vendor", ""),
          it.get("product", ""), _rem_text(it.get("remediation")),
          "\n".join(a.get("url", "") for a in (it.get("remediation") or {}).get("advisories", []))]
         for it in items])
    return _save(wb)


def build_findings_xlsx(findings: list, stats: dict = None,
                        title: str = "JATAYU — Vulnerability Findings") -> bytes:
    """Ingested connector findings workbook."""
    findings = findings or []
    stats = stats or {}
    wb = Workbook()
    by_sev = stats.get("by_severity", {}) or {}
    _summary_sheet(wb, title, [
        ("Generated", _stamp()),
        ("Total findings", stats.get("total", len(findings))),
        ("Distinct assets", stats.get("assets", "")),
        ("Unique CVEs", stats.get("unique_cves", "")),
        ("Critical", by_sev.get("CRITICAL", 0)),
        ("High", by_sev.get("HIGH", 0)),
        ("Medium", by_sev.get("MEDIUM", 0)),
        ("Low", by_sev.get("LOW", 0)),
        ("Sources", ", ".join(f"{k}: {v}" for k, v in (stats.get("by_source", {}) or {}).items())),
    ])
    _add_sheet(wb, "Findings",
        ["Source", "Asset", "IP", "Hostname", "Severity", "CVSS", "CVE(s)", "Port",
         "Plugin / Vulnerability", "First Seen"],
        [[f.get("source", ""), f.get("asset", ""), f.get("ip", ""), f.get("hostname", ""),
          f.get("severity", ""), f.get("cvss", ""), ", ".join(f.get("cve_ids") or []),
          f.get("port", ""), f.get("plugin_name", ""), f.get("first_seen", "")]
         for f in findings],
        sev_col=4)
    return _save(wb)


def build_intel_xlsx(data: dict, title: str = "JATAYU — Threat Intelligence") -> bytes:
    """Threat-intel snapshot workbook (KEV / NVD / IOCs / actors / breaches / zero-day)."""
    data = data or {}
    adv = data.get("advisories", {}) or {}
    ti = data.get("threat_intel", {}) or {}
    iocs = ti.get("iocs", {}) or {}
    dw = ti.get("darkweb", {}) or {}
    wb = Workbook()
    _summary_sheet(wb, title, [
        ("Generated", _stamp()),
        ("Top exploited CVEs (KEV)", len(adv.get("top_exploited_cves", []) or [])),
        ("Critical CVEs (NVD)", len(adv.get("immediate_attention_cves", []) or [])),
        ("Malicious IPs", len(iocs.get("malicious_ips", []) or [])),
        ("Malicious URLs", len(iocs.get("malicious_urls", []) or [])),
        ("Malware hashes", len(iocs.get("malware_hashes", []) or [])),
        ("Threat actors", len(ti.get("top_threat_actors", []) or [])),
        ("Breaches", len(dw.get("breaches", []) or [])),
        ("Zero-day records", len(ti.get("zeroday", []) or [])),
    ])
    _add_sheet(wb, "Top Exploited CVEs (KEV)",
        ["CVE", "Vendor", "Product", "Date Added", "Due Date", "Ransomware", "Description"],
        [[c.get("cve_id", ""), c.get("vendor", ""), c.get("product", ""), c.get("date_added", ""),
          c.get("due_date", ""), c.get("known_ransomware", ""), (c.get("description", "") or "")[:160]]
         for c in adv.get("top_exploited_cves", []) or []])
    _add_sheet(wb, "Critical CVEs (NVD)",
        ["CVE", "CVSS", "Ver", "Attack Vector", "Severity", "Published", "Description"],
        [[c.get("cve_id", ""), c.get("cvss_score", ""), c.get("cvss_version", ""),
          c.get("attack_vector", ""), c.get("severity", ""), c.get("published", ""),
          (c.get("description", "") or "")[:160]]
         for c in adv.get("immediate_attention_cves", []) or []])
    _add_sheet(wb, "Malicious IPs",
        ["IP", "Port", "Malware", "Country", "Status", "First Seen"],
        [[i.get("ip", ""), i.get("port", ""), i.get("malware", ""), i.get("country", ""),
          i.get("status", ""), i.get("first_seen", "")]
         for i in iocs.get("malicious_ips", []) or []])
    _add_sheet(wb, "Threat Actors",
        ["Name", "MITRE ID", "Aliases", "Country", "Sophistication"],
        [[a.get("name", ""), a.get("mitre_id", ""), "|".join(a.get("aliases") or []),
          a.get("country", ""), a.get("sophistication", "")]
         for a in ti.get("top_threat_actors", []) or []])
    _add_sheet(wb, "Breaches",
        ["Title", "Domain", "Breach Date", "Records", "Verified"],
        [[b.get("title", ""), b.get("domain", ""), b.get("breach_date", ""),
          b.get("pwn_count", 0), b.get("is_verified", "")]
         for b in dw.get("breaches", []) or []])
    _add_sheet(wb, "Zero-Day",
        ["CVE", "Title", "Type", "Software", "Actively Exploited", "Discovered"],
        [[z.get("cve", ""), z.get("title", ""), z.get("vuln_type", ""), z.get("software", ""),
          "Yes" if z.get("actively_exploited") else "No", z.get("date_discovered", "")]
         for z in ti.get("zeroday", []) or []])
    return _save(wb)


def build_remediation_xlsx(grouped: dict, title: str = "JATAYU — Remediation Plan") -> bytes:
    """Per-CVE remediation document for a CVE or a group of similar CVEs.

    `grouped` is the output of risk.group_remediations: {"groups":[{label, items:[rem]}]}.
    """
    grouped = grouped or {}
    groups = grouped.get("groups", []) or []
    wb = Workbook()
    _summary_sheet(wb, title, [
        ("Generated", _stamp()),
        ("Groups", len(groups)),
        ("Total CVEs", grouped.get("count", 0)),
    ] + [(f"  {g.get('label', '')}", len(g.get("items", []))) for g in groups[:20]])
    rows = []
    for g in groups:
        for rem in g.get("items", []) or []:
            vendor_text = rem.get("advisory_text", "")
            if vendor_text and rem.get("advisory_url"):
                vendor_text += f"\n[source: {rem['advisory_url']}]"
            rows.append([
                g.get("label", ""), rem.get("cve_id", ""), rem.get("sla", ""),
                rem.get("summary", ""), _rem_text(rem), vendor_text,
                "\n".join(a.get("url", "") for a in rem.get("advisories", []) or []),
            ])
    _add_sheet(wb, "Remediation Plan",
        ["Group", "CVE", "SLA / Due", "Summary", "Detailed Steps",
         "Vendor Advisory (live)", "Advisories"], rows)
    return _save(wb)


def build_custom_report_xlsx(sections: list, data: dict = None, risk_result: dict = None,
                             findings: list = None, stats: dict = None) -> bytes:
    """Combined workbook assembled from the chosen sections (report builder)."""
    sections = [s.lower() for s in (sections or [])]
    wb = Workbook()
    _summary_sheet(wb, "JATAYU — Custom Report", [
        ("Generated", _stamp()),
        ("Sections", ", ".join(sections) or "—"),
    ])

    def _merge(src_bytes):
        """Append all non-Summary sheets from a builder's workbook into wb."""
        from openpyxl import load_workbook
        src = load_workbook(io.BytesIO(src_bytes))
        for name in src.sheetnames:
            if name == "Summary":
                continue
            s = src[name]
            dst = wb.create_sheet(title=name[:31])
            for row in s.iter_rows(values_only=True):
                dst.append(row)
            _style_header(dst, 1, max(1, s.max_column))
            _autosize(dst)

    if "risk" in sections and risk_result is not None:
        _merge(build_risk_xlsx(risk_result))
    if "findings" in sections and findings is not None:
        _merge(build_findings_xlsx(findings, stats))
    if any(s in sections for s in ("threat-intel", "intel", "threat_intel", "zeroday")) and data is not None:
        _merge(build_intel_xlsx(data))
    return _save(wb)


# ── Generic per-section export (one shape feeds both XLSX and PDF) ──────────────
# `tables` is a list of {"title": str, "columns": [str], "rows": [[cell, ...]]}.

_ILLEGAL_XLSX = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _xlsx_safe(v):
    """openpyxl rejects control chars (binary IOC data etc.) — sanitise cells."""
    if isinstance(v, (int, float, bool)) or v is None:
        return v
    return _ILLEGAL_XLSX.sub("", str(v))


def build_section_xlsx(title: str, tables: list) -> bytes:
    """One styled sheet per table — used by every section's Export → XLSX."""
    tables = tables or []
    wb = Workbook()
    _summary_sheet(wb, title, [("Generated", _stamp()),
                               ("Sections", ", ".join(t.get("title", "") for t in tables) or "—")])
    if not tables:
        _add_sheet(wb, "Data", ["(no data)"], [])
    for t in tables:
        rows = [[_xlsx_safe(c) for c in r] for r in (t.get("rows") or [])]
        _add_sheet(wb, (t.get("title") or "Data")[:31], t.get("columns") or [], rows)
    return _save(wb)


def build_section_pdf(title: str, tables: list, subtitle: str = "") -> bytes:
    """Titled, styled multi-table PDF — used by every section's Export → PDF."""
    ss = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title=title)
    story = [Paragraph("JATAYU", ss["JTitle"]),
             Paragraph(title, ss["JSub"]),
             Paragraph(subtitle or _stamp(), ss["JSub"]), Spacer(1, 8)]
    usable = 182 * mm
    for t in (tables or []):
        cols = t.get("columns") or []
        rows = t.get("rows") or []
        story.append(Paragraph(t.get("title", "") or "Data", ss["JH2"]))
        if not cols or not rows:
            story.append(Paragraph("No data.", ss["JBody"]))
            story.append(Spacer(1, 6))
            continue
        header = [Paragraph(f"<b>{_html.escape(str(c))}</b>", ss["JCell"]) for c in cols]
        data = [header]
        for r in rows[:400]:
            # Truncate before escaping so entity sequences (&amp;, &lt;) are never split mid-token
            data.append([Paragraph(_html.escape(str(c if c is not None else "")[:240]), ss["JCell"])
                         for c in (list(r) + [""] * (len(cols) - len(r)))[:len(cols)]])
        cw = usable / max(1, len(cols))
        tbl = Table(data, colWidths=[cw] * len(cols), repeatRows=1)
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B2545")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D0D7E2")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5FA")]),
            ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]))
        story.append(tbl)
        story.append(Spacer(1, 10))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Generated by JATAYU Threat Intelligence Platform.", ss["JSub"]))
    doc.build(story)
    return buf.getvalue()


# ════════════════════════════════════════════════════════════════════════════
#  PDF
# ════════════════════════════════════════════════════════════════════════════

def _styles():
    ss = getSampleStyleSheet()
    ss.add(ParagraphStyle("JTitle", parent=ss["Title"], fontSize=22,
                          textColor=colors.HexColor("#0B2545"), spaceAfter=4))
    ss.add(ParagraphStyle("JSub", parent=ss["Normal"], fontSize=10,
                          textColor=colors.HexColor("#5673A0"), spaceAfter=2))
    ss.add(ParagraphStyle("JH2", parent=ss["Heading2"], fontSize=14,
                          textColor=colors.HexColor("#0B2545"), spaceBefore=14, spaceAfter=6))
    ss.add(ParagraphStyle("JBody", parent=ss["Normal"], fontSize=9, leading=12))
    ss.add(ParagraphStyle("JCell", parent=ss["Normal"], fontSize=8, leading=10))
    ss.add(ParagraphStyle("JCellW", parent=ss["Normal"], fontSize=8, leading=10,
                          textColor=colors.white))
    return ss


def _sev_para(sev, ss):
    return Paragraph(
        f'<font color="white"><b>{sev}</b></font>',
        ParagraphStyle("s", parent=ss["JCell"], alignment=TA_CENTER))


def build_pdf(result: dict) -> bytes:
    ss = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm,
                            title=f"Attack Surface Report — {result.get('query','')}")
    story = []
    rs = result.get("risk_score", {}) or {}
    findings = result.get("findings", []) or []
    sev_counts = {k: 0 for k in SEV_ORDER}
    for f in findings:
        sev_counts[f.get("severity", "INFO")] = sev_counts.get(f.get("severity", "INFO"), 0) + 1

    # ── Cover / header ───────────────────────────────────────────────────────
    story.append(Paragraph("JATAYU", ss["JTitle"]))
    story.append(Paragraph("External Attack Surface Report", ss["JSub"]))
    story.append(Spacer(1, 8))

    level = rs.get("level", "UNKNOWN")
    rc = colors.HexColor("#" + RISK_COLORS.get(level, "5673A0"))
    cover = Table([
        [Paragraph(f"<b>Target</b><br/>{result.get('query','')}", ss["JBody"]),
         Paragraph(f"<b>Type</b><br/>{(result.get('type','') or '').upper()}", ss["JBody"]),
         Paragraph(f"<b>Scan time</b><br/>{_fmt_ts(result.get('timestamp',''))}", ss["JBody"])],
        [Paragraph(f'<font color="white" size="20"><b>{rs.get("score",0)}/100</b></font>', ss["JBody"]),
         Paragraph(f'<font color="white"><b>RISK: {level}</b></font>', ss["JBody"]),
         Paragraph(f'<font color="white">{sum(1 for h in result.get("hosts",[]) if h.get("alive"))} live hosts · '
                   f'{len(result.get("vulnerabilities",[]))} CVEs · {len(findings)} findings</font>', ss["JBody"])],
    ], colWidths=[60 * mm, 50 * mm, 68 * mm])
    cover.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D0D7E2")),
        ("INNERGRID", (0, 0), (-1, 0), 0.5, colors.HexColor("#D0D7E2")),
        ("BACKGROUND", (0, 1), (-1, 1), rc),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(cover)
    story.append(Spacer(1, 10))

    # ── Severity summary strip ───────────────────────────────────────────────
    strip = Table([[
        Paragraph(f'<font color="white"><b>{sev_counts["CRITICAL"]}</b><br/>Critical</font>', ss["JCellW"]),
        Paragraph(f'<font color="white"><b>{sev_counts["HIGH"]}</b><br/>High</font>', ss["JCellW"]),
        Paragraph(f'<font color="white"><b>{sev_counts["MEDIUM"]}</b><br/>Medium</font>', ss["JCellW"]),
        Paragraph(f'<font color="white"><b>{sev_counts["LOW"]}</b><br/>Low</font>', ss["JCellW"]),
        Paragraph(f'<font color="white"><b>{sev_counts["INFO"]}</b><br/>Info</font>', ss["JCellW"]),
    ]], colWidths=[35.6 * mm] * 5)
    strip.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#C00000")),
        ("BACKGROUND", (1, 0), (1, 0), colors.HexColor("#E8741C")),
        ("BACKGROUND", (2, 0), (2, 0), colors.HexColor("#C9A400")),
        ("BACKGROUND", (3, 0), (3, 0), colors.HexColor("#2E8B57")),
        ("BACKGROUND", (4, 0), (4, 0), colors.HexColor("#5673A0")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(strip)

    # ── Risk factors ─────────────────────────────────────────────────────────
    if rs.get("factors"):
        story.append(Paragraph("Executive Summary", ss["JH2"]))
        for fct in rs["factors"]:
            story.append(Paragraph(f"• {fct}", ss["JBody"]))

    # ── VM / ASM expert assessment (see asm.py / ASM/SKILL.md) ────────────────
    asm_data = result.get("asm") or {}
    asm_findings = asm_data.get("findings") or []
    if asm_data:
        sm = asm_data.get("summary", {})
        cn = sm.get("counts", {})
        story.append(Paragraph("VM / ASM Expert Assessment", ss["JH2"]))
        if sm.get("key_risk_statement"):
            story.append(Paragraph(sm["key_risk_statement"], ss["JBody"]))
        story.append(Paragraph(
            f"<b>Findings:</b> {cn.get('total',0)} &nbsp;&nbsp; "
            f"<b>Critical:</b> {cn.get('critical',0)} &nbsp;&nbsp; <b>High:</b> {cn.get('high',0)} &nbsp;&nbsp; "
            f"<b>Actively Exploited (CISA KEV):</b> {cn.get('kev',0)} &nbsp;&nbsp; "
            f"<b>Ransomware-linked:</b> {cn.get('ransomware',0)}", ss["JBody"]))
        for act in (sm.get("immediate_actions") or [])[:4]:
            story.append(Paragraph(f"• {act}", ss["JBody"]))
        if asm_findings:
            story.append(Spacer(1, 4))
            data = [[Paragraph(f"<b>{c}</b>", ss["JCell"]) for c in
                     ("ID", "Sev", "Confidence", "Standards (OWASP / CWE)", "KEV", "Remediate", "Finding")]]
            style = [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B2545")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D0D7E2")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
            for i, f in enumerate(asm_findings[:25]):
                sev = f.get("severity", "INFO")
                st = f.get("standards", {})
                std = f"{st.get('owasp', '—')}<br/>{st.get('cwe', '—')}"
                data.append([
                    Paragraph(f.get("id", ""), ss["JCell"]),
                    _sev_para(sev, ss),
                    Paragraph(f.get("confidence", ""), ss["JCell"]),
                    Paragraph(std, ss["JCell"]),
                    Paragraph(st.get("kev", "No"), ss["JCell"]),
                    Paragraph(f.get("remediation_timeline", ""), ss["JCell"]),
                    Paragraph(f.get("title", ""), ss["JCell"]),
                ])
                style.append(("BACKGROUND", (1, i + 1), (1, i + 1),
                              colors.HexColor("#" + SEV_COLORS.get(sev, "5673A0"))))
            t = Table(data, colWidths=[15 * mm, 14 * mm, 18 * mm, 40 * mm, 10 * mm, 24 * mm, 57 * mm], repeatRows=1)
            t.setStyle(TableStyle(style))
            story.append(t)
        story.append(Spacer(1, 8))

    # ── Prioritised remediation ──────────────────────────────────────────────
    remed = result.get("remediation", [])
    if remed:
        story.append(Paragraph("Prioritized Remediation Roadmap", ss["JH2"]))
        data = [[Paragraph("<b>#</b>", ss["JCell"]), Paragraph("<b>Sev</b>", ss["JCell"]),
                 Paragraph("<b>Issue</b>", ss["JCell"]), Paragraph("<b>Hosts</b>", ss["JCell"]),
                 Paragraph("<b>Recommended Fix</b>", ss["JCell"])]]
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B2545")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D0D7E2")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
        for i, r in enumerate(remed[:25]):
            sev = r.get("severity", "INFO")
            data.append([
                Paragraph(str(i + 1), ss["JCell"]),
                _sev_para(sev, ss),
                Paragraph(r.get("title", ""), ss["JCell"]),
                Paragraph(str(r.get("affected_count", 0)), ss["JCell"]),
                Paragraph(r.get("remediation", ""), ss["JCell"]),
            ])
            style.append(("BACKGROUND", (1, i + 1), (1, i + 1),
                          colors.HexColor("#" + SEV_COLORS.get(sev, "5673A0"))))
        t = Table(data, colWidths=[8 * mm, 16 * mm, 58 * mm, 12 * mm, 84 * mm], repeatRows=1)
        t.setStyle(TableStyle(style))
        story.append(t)

    # ── Live host inventory ──────────────────────────────────────────────────
    live = [h for h in result.get("hosts", []) if h.get("alive")]
    if live:
        story.append(PageBreak())
        story.append(Paragraph("Live Host Inventory", ss["JH2"]))
        data = [[Paragraph(f"<b>{c}</b>", ss["JCell"]) for c in
                 ("Host", "HTTP", "IP", "Ports", "CVEs", "Findings", "Tech")]]
        for h in live[:60]:
            data.append([
                Paragraph(h.get("host", ""), ss["JCell"]),
                Paragraph(str(h.get("status", "") or ""), ss["JCell"]),
                Paragraph(", ".join(h.get("ips", [])[:2]), ss["JCell"]),
                Paragraph(str(len(h.get("ports", []))), ss["JCell"]),
                Paragraph(str(len(h.get("cves", []))), ss["JCell"]),
                Paragraph(str(len(h.get("findings", []))), ss["JCell"]),
                Paragraph(", ".join(h.get("tech", [])[:3]), ss["JCell"]),
            ])
        t = Table(data, colWidths=[52 * mm, 14 * mm, 40 * mm, 13 * mm, 13 * mm, 16 * mm, 30 * mm], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B2545")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D0D7E2")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5FA")]),
            ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]))
        story.append(t)

    # ── Detailed findings ────────────────────────────────────────────────────
    if findings:
        story.append(PageBreak())
        story.append(Paragraph("Detailed Findings", ss["JH2"]))
        data = [[Paragraph(f"<b>{c}</b>", ss["JCell"]) for c in
                 ("Sev", "Category", "Asset", "Finding", "Remediation")]]
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B2545")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D0D7E2")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]
        for i, f in enumerate(findings[:120]):
            sev = f.get("severity", "INFO")
            data.append([
                _sev_para(sev, ss),
                Paragraph(f.get("category", ""), ss["JCell"]),
                Paragraph(f.get("asset", ""), ss["JCell"]),
                Paragraph(f.get("title", ""), ss["JCell"]),
                Paragraph(f.get("remediation", ""), ss["JCell"]),
            ])
            style.append(("BACKGROUND", (0, i + 1), (0, i + 1),
                          colors.HexColor("#" + SEV_COLORS.get(sev, "5673A0"))))
        t = Table(data, colWidths=[16 * mm, 26 * mm, 40 * mm, 48 * mm, 48 * mm], repeatRows=1)
        t.setStyle(TableStyle(style))
        story.append(t)

    # ── Footer note ──────────────────────────────────────────────────────────
    story.append(Spacer(1, 12))
    story.append(Paragraph(
        "Generated by JATAYU Threat Intelligence Platform. Sources: "
        + ", ".join(result.get("sources_used", [])) + ". "
        "This report reflects passive + non-intrusive active reconnaissance of the "
        "specified target's external attack surface.",
        ss["JSub"]))

    doc.build(story)
    return buf.getvalue()


# ════════════════════════════════════════════════════════════════════════════
#  RESEARCH INTEL BRIEF (Big-4 style) — implements intel.md §5–6
# ════════════════════════════════════════════════════════════════════════════

_IOC_IPV4 = re.compile(r"\b(?:\d{1,3}\.){3}\d{1,3}\b")
_IOC_SHA  = re.compile(r"\b[a-fA-F0-9]{64}\b")
_IOC_MD5  = re.compile(r"\b[a-fA-F0-9]{32}\b")
_IOC_ETH  = re.compile(r"\b0x[a-fA-F0-9]{40}\b")
_IOC_CVE  = re.compile(r"CVE-\d{4}-\d{4,7}", re.I)
_IOC_DOM  = re.compile(r"\b(?:[a-z0-9-]+\.)+(?:com|net|org|io|ru|cn|xyz|info|biz|co|dev|gov|edu)\b", re.I)
_MD_STRIP = re.compile(r"[`*#>_]+")
_TAGS     = re.compile(r"<[^>]+>")


def _plain(text: str) -> str:
    """Markdown/HTML → plain text for PDF paragraphs."""
    if not text:
        return ""
    t = _TAGS.sub(" ", text)
    t = re.sub(r"```.*?```", " [code omitted] ", t, flags=re.S)
    t = _MD_STRIP.sub("", t)
    t = _html.unescape(t)
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def _extract_iocs(text: str) -> dict:
    blob = text or ""
    ips    = sorted({i for i in _IOC_IPV4.findall(blob)
                     if not i.startswith(("0.", "255.")) and i.count(".") == 3})[:25]
    shas   = sorted(set(_IOC_SHA.findall(blob)))[:20]
    md5s   = sorted({m for m in _IOC_MD5.findall(blob) if m not in shas})[:20]
    cves   = sorted({c.upper() for c in _IOC_CVE.findall(blob)})[:25]
    doms   = sorted(set(_IOC_DOM.findall(blob)))[:25]
    eth    = sorted(set(_IOC_ETH.findall(blob)))[:25]
    return {"ips": ips, "sha256": shas, "md5": md5s, "cves": cves, "domains": doms, "addresses": eth}


def research_report_filename(item: dict) -> str:
    base = re.sub(r"[^A-Za-z0-9]+", "_", (item.get("title", "brief") or "brief"))[:48].strip("_")
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    return f"threat_brief_{base}_{stamp}.pdf"


def build_research_pdf(item: dict) -> bytes:
    """Big-4 consulting style threat-intelligence brief for one research item."""
    ss = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm,
                            title=f"Threat Brief — {item.get('title','')}")
    story = []
    sev = (item.get("severity") or "INFO").upper()
    sc = colors.HexColor("#" + SEV_COLORS.get(sev, "5673A0"))
    body_text = item.get("content", "") or item.get("preview", "")
    iocs = _extract_iocs(body_text + " " + item.get("title", ""))

    # ── Header band ──────────────────────────────────────────────────────────
    story.append(Paragraph("JATAYU", ss["JTitle"]))
    story.append(Paragraph("Threat Intelligence Brief", ss["JSub"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph(_html.escape(item.get("title", "Untitled")), ss["JH2"]))

    meta = Table([[
        Paragraph(f'<font color="white"><b>{sev}</b></font>', ss["JBody"]),
        Paragraph(f"<b>Category</b>  {_html.escape(item.get('category','—'))}", ss["JBody"]),
        Paragraph(f"<b>Source</b>  {_html.escape(item.get('source','—'))}", ss["JBody"]),
        Paragraph(f"<b>Published</b>  {_html.escape(str(item.get('published',''))[:16] or '—')}", ss["JBody"]),
    ]], colWidths=[22 * mm, 56 * mm, 50 * mm, 50 * mm])
    meta.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D0D7E2")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D0D7E2")),
        ("BACKGROUND", (0, 0), (0, 0), sc),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(meta)
    story.append(Spacer(1, 12))

    # ── Executive summary ────────────────────────────────────────────────────
    story.append(Paragraph("Executive Summary", ss["JH2"]))
    story.append(Paragraph(_html.escape(item.get("preview", "")) or "No summary available.", ss["JBody"]))

    # ── Indicators of Compromise ─────────────────────────────────────────────
    ioc_rows = []
    label = {"cves": "CVEs", "ips": "IP Addresses", "domains": "Domains",
             "sha256": "SHA-256", "md5": "MD5", "addresses": "Wallet / Contract Addresses"}
    for key in ("cves", "ips", "domains", "sha256", "md5", "addresses"):
        vals = iocs.get(key) or []
        if vals:
            ioc_rows.append([Paragraph(f"<b>{label[key]}</b>", ss["JCell"]),
                             Paragraph(", ".join(_html.escape(v) for v in vals), ss["JCell"])])
    if ioc_rows:
        story.append(Paragraph("Indicators of Compromise", ss["JH2"]))
        t = Table(ioc_rows, colWidths=[30 * mm, 148 * mm])
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D0D7E2")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F2F5FA")),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(t)

    # ── Full analysis ────────────────────────────────────────────────────────
    story.append(Paragraph("Threat Analysis", ss["JH2"]))
    plain = _plain(body_text)
    for para in [p for p in plain.split("\n\n") if p.strip()][:60]:
        story.append(Paragraph(_html.escape(para.strip()).replace("\n", "<br/>"), ss["JBody"]))
        story.append(Spacer(1, 5))

    # ── References ───────────────────────────────────────────────────────────
    if item.get("link"):
        story.append(Paragraph("References", ss["JH2"]))
        link = _html.escape(item["link"])
        story.append(Paragraph(f'<a href="{link}">{link}</a>', ss["JBody"]))

    story.append(Spacer(1, 14))
    story.append(Paragraph(
        "Generated by JATAYU Threat Intelligence Platform — Security Research module. "
        "Classification: Internal Use Only.", ss["JSub"]))

    doc.build(story)
    return buf.getvalue()
